import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
D = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
from pdf_data import variants, wide_rows, P

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(D, "Bang_gia_mau_Gom_Vu_Gia.xlsx")
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3D8A")
EX = PatternFill("solid", fgColor="EEF2FB")
thin = Side(style="thin", color="C9D2E6")
B = Border(bottom=thin, top=thin, left=thin, right=thin)

wb = Workbook()
ws0 = wb.active
ws0.title = "Hướng dẫn"
lines = [
 ("FILE EXCEL MẪU – BẢNG GIÁ GỐM VŨ GIA", True),
 ("Nguồn: chép từ PDF “Báo giá sản phẩm đồ thờ men lam – Bộ đồ thờ Phúc Liên Hoa”. Hãy kiểm tra lại giá trước khi dùng chính thức.", False),
 ("", False),
 ("Sheet “Bảng giá theo biến thể”: mỗi dòng là một biến thể (quy cách × dòng men), có Mã cố định. Đây là định dạng ứng dụng xuất ra và nhập lại.", False),
 ("Sheet “Bảng giá 2 cột giá”: giống bố cục PDF, giá Men lam và Men lam vẽ vàng nằm ở 2 cột. Khi nhập, chọn cả 2 cột giá để ứng dụng tách thành biến thể.", False),
 ("", False),
 ("Cách điền:", True),
 ("• Mã: để trống nếu chưa có – ứng dụng tự tạo mã và giữ ổn định khi nhập lại. Đã có mã thì không đổi mã.", False),
 ("• Đơn giá: ghi số, ví dụ 1200000, 1.200.000 hoặc 1,200,000. Để trống nếu chưa có giá – ứng dụng ghi “Chưa có giá”, không coi là 0.", False),
 ("• ĐVT: chiếc, đôi, bộ, cái… Giá theo đôi thì số lượng 1 là 1 đôi.", False),
 ("• Trạng thái: “Đang dùng” hoặc “Ngừng dùng”.", False),
 ("", False),
 ("Ghi chú khi chuyển từ PDF:", True),
 ("• Ba dòng “Chân đế” trong PDF (mục Bộ đỉnh hạc, Bộ trà thờ, Bộ rượu cúng) được đổi tên thành Chân đế đỉnh / Chân đế bộ trà / Chân đế bộ rượu để không bị trùng.", False),
 ("• PDF ghi giá vàng kim = 0 cho Bóng đèn và Bộ đũa gỗ: sheet biến thể để trống giá và đặt Ngừng dùng; sheet 2 cột giữ số 0 như PDF (ứng dụng sẽ báo giá không hợp lệ).", False),
 ("• Cột VÀNG KIM trong PDF được hiểu là dòng men “Men lam vẽ vàng”.", False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=13 if i == 1 else 11, color="1F3D8A" if i == 1 else "222222")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws0.column_dimensions["A"].width = 120

def header(ws, row, cols, widths):
    for j, (h, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = B
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30

# Sheet 1: biến thể
ws = wb.create_sheet("Bảng giá theo biến thể")
ws["A1"] = "BẢNG GIÁ ĐỒ THỜ MEN LAM – BỘ PHÚC LIÊN HOA (chép từ PDF)"
ws["A1"].font = Font(name=F, bold=True, size=13, color="1F3D8A")
ws["A2"] = "Dòng tiêu đề cột ở dòng 4. Đơn giá tính bằng đồng (đ)."
ws["A2"].font = Font(name=F, italic=True, color="555555")
cols = ["Mã", "Nhóm", "Tên sản phẩm", "Họa tiết", "Quy cách", "Dòng men", "ĐVT", "Đơn giá", "Ghi chú", "Trạng thái"]
header(ws, 4, cols, [18, 18, 32, 10, 12, 18, 8, 14, 46, 12])
for i, v in enumerate(variants(), 5):
    vals = [v["code"], v["group"], v["name"], v["pattern"], v["size"], v["glaze"], v["unit"], v["price"], v["note"],
            "Đang dùng" if v["active"] else "Ngừng dùng"]
    for j, val in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=val)
        c.font = Font(name=F)
        c.border = B
        if j == 8:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        if j in (4, 5, 7, 10):
            c.alignment = Alignment(horizontal="center")
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:J{4 + len(variants())}"

# Sheet 2: 2 cột giá như PDF
ws2 = wb.create_sheet("Bảng giá 2 cột giá")
cols2 = ["STT", "Nhóm", "Tên sản phẩm", "Họa tiết", "Quy cách", "ĐVT", "Giá Men lam", "Giá Men lam vẽ vàng", "Ghi chú"]
header(ws2, 1, cols2, [6, 18, 32, 10, 12, 8, 18, 20, 46])
stt = 0
for i, r in enumerate(wide_rows(), 2):
    stt += 1
    vals = [stt, r["group"], r["name"], r["pattern"], r["size"], r["unit"], r["ml"], r["vk"], r["note"]]
    for j, val in enumerate(vals, 1):
        c = ws2.cell(row=i, column=j, value=val)
        c.font = Font(name=F)
        c.border = B
        if j in (7, 8):
            c.number_format = "#,##0"
    ws2.cell(row=i, column=1).alignment = Alignment(horizontal="center")
ws2.freeze_panes = "A2"

wb.active = 1
wb.save(OUT)

# JSON nhúng vào ứng dụng
with open(os.path.join(D, "pdf_data.json"), "w", encoding="utf-8", newline="\n") as f:
    json.dump({"variants": variants(), "wide": wide_rows()}, f, ensure_ascii=False)
print("ok", OUT, len(variants()))
